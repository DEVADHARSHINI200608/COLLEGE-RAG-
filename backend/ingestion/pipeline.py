"""
Document ingestion pipeline — parse → chunk → embed → store.
Modular: add new parsers by registering them in PARSER_MAP.
"""
from __future__ import annotations

import hashlib
import io
import os
from pathlib import Path
from typing import List, Optional, Tuple

import structlog
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from config import settings
from retrieval.vector_store import add_documents

logger = structlog.get_logger(__name__)


# ─── Text Splitter ──────────────────────────────────────────────────────
def _get_splitter() -> RecursiveCharacterTextSplitter:
    return RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        length_function=len,
        separators=["\n\n", "\n", " ", ""],
    )


# ─── Parsers ────────────────────────────────────────────────────────────
def parse_pdf(file_bytes: bytes, filename: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n\n".join(
        page.extract_text() or "" for page in reader.pages
    )


def parse_docx(file_bytes: bytes, filename: str) -> str:
    from docx import Document as DocxDocument
    doc = DocxDocument(io.BytesIO(file_bytes))
    return "\n\n".join(para.text for para in doc.paragraphs if para.text.strip())


def parse_pptx(file_bytes: bytes, filename: str) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(file_bytes))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
    return "\n\n".join(texts)


def parse_txt(file_bytes: bytes, filename: str) -> str:
    return file_bytes.decode("utf-8", errors="replace")


def parse_markdown(file_bytes: bytes, filename: str) -> str:
    return file_bytes.decode("utf-8", errors="replace")


def parse_csv(file_bytes: bytes, filename: str) -> str:
    """For RAG purposes — return as text. Structured queries use the raw CSV."""
    import pandas as pd
    df = pd.read_csv(io.BytesIO(file_bytes))
    return df.to_string(index=False)


def parse_xlsx(file_bytes: bytes, filename: str) -> str:
    import pandas as pd
    df = pd.read_excel(io.BytesIO(file_bytes))
    return df.to_string(index=False)


# ─── Parser registry ────────────────────────────────────────────────────
PARSER_MAP = {
    ".pdf": parse_pdf,
    ".doc": parse_docx,
    ".docx": parse_docx,
    ".ppt": parse_pptx,
    ".pptx": parse_pptx,
    ".txt": parse_txt,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".csv": parse_csv,
    ".xlsx": parse_xlsx,
    ".xls": parse_xlsx,
}

SUPPORTED_EXTENSIONS = set(PARSER_MAP.keys())


def validate_file(filename: str, file_bytes: bytes) -> Tuple[bool, str]:
    """Validate file type and size."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        return False, f"Unsupported file type: {ext}"
    if len(file_bytes) > settings.max_file_size_bytes:
        return False, f"File too large (max {settings.MAX_FILE_SIZE_MB}MB)"
    return True, "OK"


def compute_hash(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


def parse_file(filename: str, file_bytes: bytes) -> str:
    """Dispatch to the correct parser."""
    ext = Path(filename).suffix.lower()
    parser = PARSER_MAP.get(ext)
    if not parser:
        raise ValueError(f"No parser for extension: {ext}")
    return parser(file_bytes, filename)


def chunk_text(text: str, metadata: dict) -> List[Document]:
    """Split text into chunks, preserving metadata."""
    splitter = _get_splitter()
    return splitter.create_documents([text], metadatas=[metadata])


async def ingest_file(
    filename: str,
    file_bytes: bytes,
    source_id: str,
    source_name: str,
    classification: str,
    access_roles: List[str],
    allowed_users: List[str],
    allowed_departments: List[str],
    extra_metadata: Optional[dict] = None,
) -> Tuple[int, str]:
    """
    Full ingestion pipeline:
    file_bytes → parse → chunk → embed → vector store.

    Returns (chunk_count, content_hash).
    """
    logger.info("ingestion_started", filename=filename, source_id=source_id)

    # Validate
    ok, msg = validate_file(filename, file_bytes)
    if not ok:
        raise ValueError(msg)

    # Hash for deduplication
    content_hash = compute_hash(file_bytes)

    # Parse
    try:
        text = parse_file(filename, file_bytes)
    except Exception as exc:
        logger.error("parse_failed", filename=filename, error=str(exc))
        raise ValueError(f"Failed to parse {filename}: {exc}")

    if not text.strip():
        raise ValueError(f"No text content extracted from {filename}")

    # Chunk
    base_metadata = {
        "filename": filename,
        "source_id": source_id,
        "source_name": source_name,
        **(extra_metadata or {}),
    }
    chunks = chunk_text(text, base_metadata)

    # Embed + store
    chunk_count = await add_documents(
        documents=chunks,
        classification=classification,
        source_id=source_id,
        source_name=source_name,
        access_roles=access_roles,
        allowed_users=allowed_users,
        allowed_departments=allowed_departments,
    )

    logger.info(
        "ingestion_complete",
        filename=filename,
        source_id=source_id,
        chunks=chunk_count,
    )
    return chunk_count, content_hash
